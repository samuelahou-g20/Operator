"""Fill the FROZEN G-20 sentiment-brief template with a day's content, and convert to PDF.

The format is locked in ``templates/g20_sentiment_brief.docx`` (the colleagues' approved version,
version-controlled). This module does NOT build layout — it only replaces the text/number content
of the template's known slots, preserving every style/format from the file, so the look can never
drift in code. ``to_pdf`` renders the filled .docx via LibreOffice for an exact-match client PDF.

Slot map (indices are fixed because the template is frozen; a structure guard fails loudly if the
template is edited): para 2 = abstract, 6 = net line, 10 = trending (run 2), 13-15 = voices,
18 = takeaway (run 1); table 0 = 5-cell KPI strip; table 1 = 8 theme rows.

MUST NOT: read the lake, fetch data, or restyle the template — content injection only.
"""

from __future__ import annotations

import os
import subprocess
import time

from docx import Document
from docx.shared import RGBColor

from src.common import paths

TEMPLATE = str(paths.REPO_ROOT / "templates" / "g20_sentiment_brief.docx")
NAVY, RED, GRAY = "1B2A4A", "B5342B", "8A8F98"


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _score_hex(v: float) -> str:
    return NAVY if v >= 0.05 else RED if v <= -0.05 else GRAY


def fill(report: dict, out_docx: str, template: str = TEMPLATE) -> str:
    """Inject `report` into the frozen template, preserving its formatting. Returns out_docx."""
    d = Document(template)
    t = d.tables
    if len(t) != 2 or len(t[0].columns) != 5 or len(t[1].rows) != 8 or len(d.paragraphs) < 19:
        raise ValueError("frozen template structure changed — refusing to fill (re-map slots)")
    p = d.paragraphs

    p[2].runs[0].text = report["abstract"]
    p[6].runs[0].text = report["net_line"]
    p[10].runs[2].text = report["trending"]                 # after "TRENDING" + spaces

    voices = [v.strip().rstrip(";") for v in report["voices"].split(";") if v.strip()]
    lines = voices[:2] + (["; ".join(voices[2:])] if len(voices) > 2 else [])
    for i, idx in enumerate((13, 14, 15)):
        p[idx].runs[0].text = lines[i] if i < len(lines) else ""

    tk = report["takeaway"]
    p[18].runs[1].text = " " + (tk.split(":", 1)[1].strip() if ":" in tk else tk)

    for cell, (val, lab, _c) in zip(t[0].rows[0].cells, report["kpis"], strict=True):
        cell.paragraphs[0].runs[0].text = val
        cell.paragraphs[0].runs[0].font.color.rgb = _rgb(_c)
        cell.paragraphs[1].runs[0].text = lab.upper()

    for row, th in zip(t[1].rows, report["themes"], strict=True):
        row.cells[0].paragraphs[0].runs[0].text = th["name"]
        sc = row.cells[1].paragraphs[0].runs[0]
        sc.text = f"{th['score']:+.2f}"
        sc.font.color.rgb = _rgb(_score_hex(th["score"]))
        c2 = row.cells[2].paragraphs[0]
        c2.runs[0].text = th["line"] + " "
        if len(c2.runs) > 1:
            c2.runs[1].text = th.get("contributors", "")

    d.save(out_docx)
    return out_docx


def to_pdf(docx_path: str, out_dir: str) -> str:
    """Render the filled .docx to PDF via headless LibreOffice (exact match of the template).

    Kills any running soffice first: a lingering headless instance makes ``--convert-to`` hand off
    and return BEFORE the file is written (the stale-file race that produced false page counts).
    With no instance running the conversion is synchronous; we then assert the PDF is fresh."""
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    subprocess.run(["pkill", "-f", "soffice"], capture_output=True)
    time.sleep(1.5)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
        check=True, capture_output=True, timeout=180,
    )
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(docx_path):
        raise RuntimeError("soffice did not produce a fresh PDF")
    return out
