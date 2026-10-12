"""G-20 house report theme as an editable .docx (python-docx) — for colleagues to fix format.

Mirrors src/reporting/g20.py (HTML/PDF theme) but writes a Word document from the same
structured ``report`` dict, so a draft can be edited/re-branded in Word before we lock the
final look back into the PDF path. Uses portable fonts (Georgia / Arial) so it renders on any
Word install.

MUST NOT: read files/lake, fetch data, or embed a specific report's content — the caller supplies
the ``report`` dict.
"""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY, RED, AMBER, GRAY = "1b2a4a", "b5342b", "b0902e", "8a8f98"
TAN, RULE = "efe6d3", "d8dce2"
SERIF, SANS = "Georgia", "Arial"


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _run(p, text, *, font=SERIF, size=10.5, bold=False, italic=False, color=None,
         caps=False, spacing=None):
    r = p.add_run(text.upper() if caps else text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = _rgb(color)
    # ensure the font applies to all script ranges
    r._element.rPr.rFonts.set(qn("w:cs"), font)
    if spacing is not None:                       # letter-spacing, in twentieths of a point
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        r._element.rPr.append(sp)
    return r


def _shade(cell, hexstr: str) -> None:
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexstr.upper())
    cell._tc.get_or_add_tcPr().append(sh)


def _borders(el, edges, color=RULE, sz=6) -> None:
    """Add borders (edges: subset of top/left/bottom/right/insideH/insideV) to a tbl or tc pr."""
    pr = el._tc.get_or_add_tcPr() if hasattr(el, "_tc") else el._tbl.tblPr
    tag = "w:tcBorders" if hasattr(el, "_tc") else "w:tblBorders"
    b = OxmlElement(tag)
    for edge in edges:
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color.upper())
        b.append(e)
    pr.append(b)


def _para_rule(p, top=True, bottom=True, color="141414", sz=6) -> None:
    pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge, on in (("top", top), ("bottom", bottom)):
        if on:
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "6")
            e.set(qn("w:color"), color)
            pbdr.append(e)
    pr.append(pbdr)


def _page_field(p) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r.font.name = SANS
    r.font.size = Pt(7.5)
    r.font.color.rgb = _rgb(GRAY)
    r._element.append(fld_begin)
    r._element.append(instr)
    r._element.append(fld_end)


def _header_footer(section, running_title: str) -> None:
    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(section.page_width - section.left_margin - section.right_margin,
                                               WD_TAB_ALIGNMENT.RIGHT)
    _run(hp, "G—20", font=SANS, size=11, bold=True, color=NAVY, spacing=1.5)
    _run(hp, "\t")
    _run(hp, running_title, font=SANS, size=8, bold=True, color=GRAY, caps=True, spacing=1.5)
    _para_rule(hp, top=False, bottom=True, color=NAVY, sz=8)

    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(section.page_width - section.left_margin - section.right_margin,
                                               WD_TAB_ALIGNMENT.RIGHT)
    _para_rule(fp, top=True, bottom=False, color=RULE, sz=4)
    _run(fp, "G-20 Research · Institutional · Confidential", font=SANS, size=7.5, color=GRAY)
    _run(fp, "\t")
    _page_field(fp)


def build_docx(report: dict, out_path: str) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(2)
    normal.line_spacing = 1.05
    _header_footer(sec, report["running_title"])

    # masthead
    _run(doc.add_paragraph(), report["kicker"], font=SANS, size=7.5, bold=True, color=GRAY,
         caps=True, spacing=3)
    _run(doc.add_paragraph(), report["title"], font=SERIF, size=22, bold=True, color=NAVY)
    _run(doc.add_paragraph(), report["subtitle"], font=SERIF, size=12.5, italic=True, color=NAVY)
    _run(doc.add_paragraph(), report["abstract"], font=SERIF, size=9, italic=True, color="333333")

    # KPI strip
    kpis = report["kpis"]
    t = doc.add_table(rows=1, cols=len(kpis))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t, ("top", "bottom", "left", "right", "insideV"), color=NAVY, sz=6)
    for cell, (val, lab, color) in zip(t.rows[0].cells, kpis, strict=True):
        _run(cell.paragraphs[0], val, font=SERIF, size=12.5, bold=True, color=color)
        _run(cell.add_paragraph(), lab, font=SANS, size=6.3, color=GRAY, caps=True, spacing=.6)

    _section_head(doc, "Net mood", report["net_deck"])
    _run(doc.add_paragraph(), report["net_line"], font=SERIF, size=9)

    _section_head(doc, "Sentiment by theme",
                  "Engagement-weighted score in [−1, +1]; contributors tagged.")
    tt = doc.add_table(rows=len(report["themes"]), cols=3)
    for i, th in enumerate(report["themes"]):
        cells = tt.rows[i].cells
        for c in cells:
            _borders(c, ("bottom",), color=RULE, sz=4)
        _run(cells[0].paragraphs[0], th["name"], font=SANS, size=8.7, bold=True, color=NAVY)
        _run(cells[1].paragraphs[0], f"{th['score']:+.2f}", font=SANS, size=8.7, bold=True,
             color=(NAVY if th["score"] >= 0.05 else RED if th["score"] <= -0.05 else GRAY))
        cells[1].paragraphs[0].alignment = 2  # right
        p = cells[2].paragraphs[0]
        _run(p, th["line"] + " ", font=SERIF, size=8.7)
        _run(p, th.get("contributors", ""), font=SERIF, size=8.7, italic=True, color=GRAY)

    if report.get("trending"):
        cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
        _shade(cell, TAN)
        p = cell.paragraphs[0]
        _run(p, "TRENDING   ", font=SANS, size=7.5, bold=True, color=NAVY, spacing=.8)
        _run(p, report["trending"], font=SERIF, size=8.5)

    if report.get("voices"):
        _section_head(doc, "Notable voices", report["voices_deck"])
        _run(doc.add_paragraph(), report["voices"], font=SERIF, size=9)

    if report.get("takeaway"):
        cp = doc.add_paragraph()
        _para_rule(cp, top=True, bottom=True, color="141414", sz=6)
        _run(cp, report["takeaway"], font=SERIF, size=9, italic=True, color=NAVY)

    doc.save(out_path)


def _section_head(doc, title: str, deck: str) -> None:
    _run(doc.add_paragraph(), title, font=SANS, size=12, bold=True, color=NAVY)
    _run(doc.add_paragraph(), deck, font=SERIF, size=9, italic=True, color=GRAY)
